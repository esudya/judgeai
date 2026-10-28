import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import httpx
from fastapi import FastAPI, File, UploadFile

from src.services.file_reader import save_upload_file, settings
from tests_uploads.test_docx_guard import archive

app = FastAPI()


@app.post("/upload")
async def upload(file: UploadFile = File(...)):
    path, size = await save_upload_file(file)
    return {"size": size, "name": path.name}


class UploadAdmissionTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.directory = Path(self.folder.name)
        self.storage = patch.object(settings, "UPLOAD_DIR", self.folder.name)
        self.storage.start()
        self.client = httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://test"
        )

    async def asyncTearDown(self):
        await self.client.aclose()
        self.storage.stop()
        self.folder.cleanup()

    async def test_renamed_docx_returns_415_and_removes_temporary_file(self):
        response = await self.client.post(
            "/upload", files={"file": ("fake.docx", b"not a docx", "application/octet-stream")}
        )
        self.assertEqual(response.status_code, 415)
        self.assertEqual(list(self.directory.iterdir()), [])
        self.assertNotIn("not a docx", response.text)

    async def test_valid_docx_preserves_original_bytes_even_when_mime_is_generic(self):
        raw = archive()
        response = await self.client.post(
            "/upload", files={"file": ("safe.docx", raw, "application/octet-stream")}
        )
        self.assertEqual(response.status_code, 200)
        saved = self.directory / response.json()["name"]
        self.assertEqual(saved.read_bytes(), raw)
        self.assertEqual(response.json()["size"], len(raw))

    async def test_docx_budget_maps_to_413_and_cleans_storage(self):
        raw = archive(extra=[("word/media/bomb.bin", b"0" * 200000)])
        response = await self.client.post(
            "/upload",
            files={
                "file": (
                    "bomb.docx",
                    raw,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(response.status_code, 413)
        self.assertEqual(list(self.directory.iterdir()), [])

    async def test_text_upload_remains_unchanged(self):
        raw = "Синтетический текст".encode()
        path, size = await save_upload_file(UploadFile(io.BytesIO(raw), filename="case.txt"))
        self.assertEqual(path.read_bytes(), raw)
        self.assertEqual(size, len(raw))

    async def test_stream_failure_does_not_leave_partial_file(self):
        class Broken:
            filename = "case.docx"

            def __init__(self):
                self.calls = 0

            async def read(self, size):
                self.calls += 1
                if self.calls == 1:
                    return b"partial bytes"
                raise OSError("synthetic read failure")

        with self.assertRaises(OSError):
            await save_upload_file(Broken())
        self.assertEqual(list(self.directory.iterdir()), [])

    async def test_real_python_docx_paragraphs_and_tables_are_accepted(self):
        import docx

        document = docx.Document()
        document.add_paragraph("Synthetic compatibility fixture")
        document.add_table(rows=2, cols=2).cell(0, 0).text = "Synthetic table"
        source = io.BytesIO()
        document.save(source)
        raw = source.getvalue()
        response = await self.client.post(
            "/upload",
            files={
                "file": (
                    "generated.DOCX",
                    raw,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual((self.directory / response.json()["name"]).read_bytes(), raw)

    async def test_stream_size_cap_still_applies_before_docx_validation(self):
        with patch.object(settings, "UPLOAD_MAX_MB", 1):
            response = await self.client.post(
                "/upload",
                files={
                    "file": ("large.docx", b"x" * (1024 * 1024 + 1), "application/octet-stream")
                },
            )
        self.assertEqual(response.status_code, 413)
        self.assertEqual(list(self.directory.iterdir()), [])
